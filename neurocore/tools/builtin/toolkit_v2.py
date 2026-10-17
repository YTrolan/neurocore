"""toolkit_v2 — дополнительные инструменты."""
import csv
import difflib
import hashlib
import json
import re
import socket
import subprocess
import sys
import urllib.request
from collections import Counter
from datetime import datetime
from pathlib import Path

from ..base import Tool

# --- опциональные зависимости (падают в None, если не установлены) ---
try:
    from docx import Document
except ImportError:
    Document = None

try:
    from pypdf import PdfReader
except ImportError:
    try:
        from PyPDF2 import PdfReader
    except ImportError:
        PdfReader = None

try:
    from openpyxl import load_workbook
except ImportError:
    load_workbook = None

try:
    import yaml
except ImportError:
    yaml = None

try:
    from PIL import Image
except ImportError:
    Image = None

try:
    import psutil
except ImportError:
    psutil = None

def _size_h(n):
    for u in ("B", "KB", "MB", "GB"):
        if n < 1024:
            return f"{n:.1f} {u}"
        n /= 1024
    return f"{n:.1f} TB"


class ReadDocxTool(Tool):
    name = "read_docx"
    description = "Прочитать текст из Word (.docx)."
    parameters = {"path": "путь"}

    def run(self, path):
        if Document is None:
            return "[error] pip install python-docx"
        try:
            d = Document(path)
            out = [p.text for p in d.paragraphs if p.text.strip()]
            for t in d.tables:
                for row in t.rows:
                    out.append(" | ".join(c.text for c in row.cells))
            return "\n".join(out) or "(пусто)"
        except Exception as e:
            return f"[error] {type(e).__name__}: {e}"


class ReadPdfTool(Tool):
    name = "read_pdf"
    description = "Извлечь текст из PDF."
    parameters = {"path": "PDF", "max_pages": "макс. страниц"}

    def run(self, path, max_pages=20):
        if PdfReader is None:
            return "[error] pip install pypdf"
        try:
            r = PdfReader(path)
            out = []
            for i, pg in enumerate(r.pages[:int(max_pages)]):
                out.append(f"--- Стр {i+1} ---\n{pg.extract_text() or ''}")
            return "\n\n".join(out) or "(пусто)"
        except Exception as e:
            return f"[error] {type(e).__name__}: {e}"


class ReadXlsxTool(Tool):
    name = "read_xlsx"
    description = "Прочитать Excel."
    parameters = {"path": "файл"}

    def run(self, path):
        if load_workbook is None:
            return "[error] pip install openpyxl"
        try:
            wb = load_workbook(path, data_only=True)
            out = []
            for sn in wb.sheetnames:
                ws = wb[sn]
                out.append(f"--- {sn} ({ws.max_row}x{ws.max_column}) ---")
                for row in ws.iter_rows(values_only=True):
                    out.append(" | ".join("" if v is None else str(v) for v in row))
            return "\n".join(out[:300])
        except Exception as e:
            return f"[error] {type(e).__name__}: {e}"


class FileTreeTool(Tool):
    name = "file_tree"
    description = "Дерево директории."
    parameters = {"path": "корень", "max_depth": "глубина"}

    def run(self, path=".", max_depth=3):
        try:
            root = Path(path)
            md = int(max_depth)
            out = [str(root.resolve()) + "/"]

            def walk(p, pref="", depth=0):
                if depth > md:
                    return
                try:
                    items = sorted(p.iterdir(),
                                   key=lambda x: (x.is_file(), x.name.lower()))
                except Exception:
                    return
                for i, it in enumerate(items):
                    if it.name in ("__pycache__", ".git", ".venv", "node_modules"):
                        continue
                    last = i == len(items) - 1
                    out.append(pref + ("└── " if last else "├── ") +
                               it.name + ("/" if it.is_dir() else ""))
                    if it.is_dir() and depth < md:
                        walk(it, pref + ("    " if last else "│   "), depth + 1)

            walk(root)
            return "\n".join(out[:500])
        except Exception as e:
            return f"[error] {type(e).__name__}: {e}"


class CountFilesTool(Tool):
    name = "count_files"
    description = "Файлы по расширениям."
    parameters = {"path": "папка"}

    def run(self, path="."):
        try:
            c = Counter()
            total = 0
            for p in Path(path).rglob("*"):
                if "__pycache__" in str(p):
                    continue
                if p.is_file():
                    c[p.suffix.lower() or "(без)"] += 1
                    total += 1
            out = [f"всего: {total}", ""]
            for ext, cnt in c.most_common(30):
                out.append(f"{cnt:>6}  {ext}")
            return "\n".join(out)
        except Exception as e:
            return f"[error] {type(e).__name__}: {e}"


class DirSizesTool(Tool):
    name = "dir_sizes"
    description = "Размеры содержимого папки."
    parameters = {"path": "папка"}

    def run(self, path="."):
        try:
            lines = []
            for it in sorted(Path(path).iterdir()):
                if it.is_dir():
                    total = sum(f.stat().st_size for f in it.rglob("*") if f.is_file())
                    lines.append(f"{_size_h(total):>10}  {it.name}/")
                else:
                    lines.append(f"{_size_h(it.stat().st_size):>10}  {it.name}")
            return "\n".join(lines[:200])
        except Exception as e:
            return f"[error] {type(e).__name__}: {e}"


class FileInfoTool(Tool):
    name = "file_info"
    description = "Размер, даты, SHA256."
    parameters = {"path": "файл"}

    def run(self, path):
        try:
            p = Path(path)
            st = p.stat()
            h = hashlib.sha256()
            with open(p, "rb") as f:
                for c in iter(lambda: f.read(65536), b""):
                    h.update(c)
            return (f"путь: {p.resolve()}\n"
                    f"размер: {_size_h(st.st_size)} ({st.st_size} байт)\n"
                    f"изменён: {datetime.fromtimestamp(st.st_mtime)}\n"
                    f"SHA256: {h.hexdigest()}")
        except Exception as e:
            return f"[error] {type(e).__name__}: {e}"


class TextStatsTool(Tool):
    name = "text_stats"
    description = "Статистика текста."
    parameters = {"path": "файл"}

    def run(self, path):
        try:
            t = Path(path).read_text(encoding="utf-8", errors="ignore")
            lines = t.split("\n")
            words = t.split()
            return (f"строк: {len(lines)}\nслов: {len(words)}\n"
                    f"символов: {len(t)}\n"
                    f"уникальных слов: {len(set(w.lower() for w in words))}")
        except Exception as e:
            return f"[error] {type(e).__name__}: {e}"


class TextReplaceFileTool(Tool):
    name = "text_replace_file"
    description = "Заменить текст в файле (regex опц.)."
    parameters = {"path": "файл", "pattern": "что искать",
                  "replacement": "чем", "regex": "true/false"}

    def run(self, path, pattern, replacement, regex="false"):
        try:
            p = Path(path)
            t = p.read_text(encoding="utf-8")
            if regex.lower() == "true":
                t2, n = re.subn(pattern, replacement, t)
            else:
                n = t.count(pattern)
                t2 = t.replace(pattern, replacement)
            p.write_text(t2, encoding="utf-8")
            return f"[ok] заменено {n} раз"
        except Exception as e:
            return f"[error] {type(e).__name__}: {e}"


class TextDiffTool(Tool):
    name = "text_diff"
    description = "Различия двух файлов."
    parameters = {"a": "файл 1", "b": "файл 2"}

    def run(self, a, b):
        try:
            ta = Path(a).read_text(encoding="utf-8", errors="ignore").splitlines()
            tb = Path(b).read_text(encoding="utf-8", errors="ignore").splitlines()
            d = difflib.unified_diff(ta, tb, lineterm="", fromfile=a, tofile=b)
            return "\n".join(d) or "(идентичны)"
        except Exception as e:
            return f"[error] {type(e).__name__}: {e}"


class JsonValidateTool(Tool):
    name = "json_validate"
    description = "Проверить JSON."
    parameters = {"path": "файл"}

    def run(self, path):
        try:
            json.loads(Path(path).read_text(encoding="utf-8"))
            return "[ok] валидный JSON"
        except json.JSONDecodeError as e:
            return f"[error] строка {e.lineno}, поз {e.colno}: {e.msg}"
        except Exception as e:
            return f"[error] {type(e).__name__}: {e}"


class CsvToJsonTool(Tool):
    name = "csv_to_json"
    description = "CSV → JSON."
    parameters = {"path": "CSV", "dst": "куда (или - для вывода)"}

    def run(self, path, dst="-"):
        try:
            rows = []
            with open(path, "r", encoding="utf-8", newline="") as f:
                for r in csv.DictReader(f):
                    rows.append(r)
            t = json.dumps(rows, ensure_ascii=False, indent=2)
            if dst == "-":
                return t[:5000]
            Path(dst).write_text(t, encoding="utf-8")
            return f"[ok] {len(rows)} строк → {dst}"
        except Exception as e:
            return f"[error] {type(e).__name__}: {e}"


class JsonToCsvTool(Tool):
    name = "json_to_csv"
    description = "JSON-массив → CSV."
    parameters = {"path": "JSON", "dst": "куда"}

    def run(self, path, dst):
        try:
            data = json.loads(Path(path).read_text(encoding="utf-8"))
            if not isinstance(data, list) or not data:
                return "[error] нужен массив"
            keys = list(data[0].keys())
            with open(dst, "w", encoding="utf-8", newline="") as f:
                w = csv.DictWriter(f, fieldnames=keys)
                w.writeheader()
                for r in data:
                    w.writerow({k: r.get(k, "") for k in keys})
            return f"[ok] {len(data)} строк → {dst}"
        except Exception as e:
            return f"[error] {type(e).__name__}: {e}"


class YamlToJsonTool(Tool):
    name = "yaml_to_json"
    description = "YAML → JSON."
    parameters = {"path": "YAML", "dst": "куда"}

    def run(self, path, dst):
        if yaml is None:
            return "[error] pip install pyyaml"
        try:
            d = yaml.safe_load(Path(path).read_text(encoding="utf-8"))
            Path(dst).write_text(json.dumps(d, ensure_ascii=False, indent=2),
                                 encoding="utf-8")
            return f"[ok] {dst}"
        except Exception as e:
            return f"[error] {type(e).__name__}: {e}"


class ImageInfoTool(Tool):
    name = "image_info"
    description = "Информация об изображении."
    parameters = {"path": "файл"}

    def run(self, path):
        if Image is None:
            return "[error] pip install pillow"
        try:
            im = Image.open(path)
            return (f"format: {im.format}\nsize: {im.width}x{im.height}\n"
                    f"mode: {im.mode}\nfile: {_size_h(Path(path).stat().st_size)}")
        except Exception as e:
            return f"[error] {type(e).__name__}: {e}"


class ImageResizeTool(Tool):
    name = "image_resize"
    description = "Изменить размер."
    parameters = {"path": "исходник", "dst": "куда", "width": "ширина"}

    def run(self, path, dst, width):
        if Image is None:
            return "[error] pip install pillow"
        try:
            im = Image.open(path)
            w = int(width)
            h = int(im.height * w / im.width)
            im.resize((w, h), Image.LANCZOS).save(dst)
            return f"[ok] {w}x{h} → {dst}"
        except Exception as e:
            return f"[error] {type(e).__name__}: {e}"


class ImageConvertTool(Tool):
    name = "image_convert"
    description = "Конвертировать формат."
    parameters = {"path": "исходник", "dst": "куда (расширение=формат)"}

    def run(self, path, dst):
        if Image is None:
            return "[error] pip install pillow"
        try:
            im = Image.open(path)
            fmt = Path(dst).suffix.lstrip(".").upper()
            if fmt == "JPG":
                fmt = "JPEG"
                if im.mode in ("RGBA", "P"):
                    im = im.convert("RGB")
            im.save(dst, fmt)
            return f"[ok] {dst}"
        except Exception as e:
            return f"[error] {type(e).__name__}: {e}"


class PublicIpTool(Tool):
    name = "public_ip"
    description = "Внешний IP."
    parameters = {}

    def run(self):
        try:
            req = urllib.request.Request("https://api.ipify.org?format=json",
                                          headers={"User-Agent": "Mozilla/5.0"})
            with urllib.request.urlopen(req, timeout=8) as r:
                return f"IP: {json.loads(r.read().decode()).get('ip', '?')}"
        except Exception as e:
            return f"[error] {type(e).__name__}: {e}"


class PortCheckTool(Tool):
    name = "port_check"
    description = "Открыт ли порт."
    parameters = {"host": "хост", "port": "порт"}

    def run(self, host, port):
        try:
            s = socket.socket(socket.AF_INET, socket.SOCK_STREAM)
            s.settimeout(3)
            r = s.connect_ex((host, int(port)))
            s.close()
            return f"{host}:{port} — {'открыт' if r == 0 else 'закрыт'}"
        except Exception as e:
            return f"[error] {type(e).__name__}: {e}"


class PingHostTool(Tool):
    name = "ping_host"
    description = "Ping хоста."
    parameters = {"host": "хост", "count": "пакетов (3)"}

    def run(self, host, count=3):
        try:
            cmd = ["ping", "-n" if sys.platform == "win32" else "-c",
                   str(int(count)), host]
            r = subprocess.run(cmd, capture_output=True, text=True,
                                timeout=int(count) * 5)
            return (r.stdout or r.stderr)[:2000] or "(пусто)"
        except Exception as e:
            return f"[error] {type(e).__name__}: {e}"


class ProcessTopTool(Tool):
    name = "process_top"
    description = "Топ процессов по CPU."
    parameters = {"n": "сколько (10)"}

    def run(self, n=10):
        if psutil is None:
            return "[error] pip install psutil"
        try:
            procs = []
            for p in psutil.process_iter(["pid", "name", "cpu_percent"]):
                try:
                    procs.append(p.info)
                except Exception:
                    pass
            procs.sort(key=lambda x: x.get("cpu_percent") or 0, reverse=True)
            return "\n".join(f"{p['pid']:>6}  {p.get('cpu_percent',0):>5.1f}%  {p['name']}"
                             for p in procs[:int(n)])
        except Exception as e:
            return f"[error] {type(e).__name__}: {e}"


class DisplayInfoTool(Tool):
    name = "display_info"
    description = "Размер экрана."
    parameters = {}

    def run(self):
        try:
            import pyautogui

            w, h = pyautogui.size()
            return f"Экран: {w}x{h}"
        except Exception as e:
            return f"[error] {type(e).__name__}: {e}"


ALL_V2_TOOLS = [
    ReadDocxTool, ReadPdfTool, ReadXlsxTool, FileTreeTool,
    CountFilesTool, DirSizesTool, FileInfoTool,
    TextStatsTool, TextReplaceFileTool, TextDiffTool,
    JsonValidateTool, CsvToJsonTool, JsonToCsvTool, YamlToJsonTool,
    ImageInfoTool, ImageResizeTool, ImageConvertTool,
    PublicIpTool, PortCheckTool, PingHostTool, ProcessTopTool,
    DisplayInfoTool,
]
